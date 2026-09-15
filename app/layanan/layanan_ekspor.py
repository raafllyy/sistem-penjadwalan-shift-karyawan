import csv
from collections.abc import Sequence
from datetime import date
from io import BytesIO, StringIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from app.layanan.layanan_jadwal import (
    buat_rentang_tanggal,
    hitung_shift,
)
from app.model.karyawan import Karyawan


def buat_matriks_jadwal(
    daftar_karyawan: Sequence[Karyawan],
    tanggal_mulai: date,
    tanggal_selesai: date,
) -> tuple[list[str], list[list[str]]]:
    """Membuat header dan baris jadwal untuk kebutuhan ekspor."""

    rentang_tanggal = buat_rentang_tanggal(
        tanggal_mulai,
        tanggal_selesai,
    )

    header = [
        "ID",
        "Nama",
        *[tanggal.strftime("%Y/%m/%d") for tanggal in rentang_tanggal],
    ]

    baris = []

    for karyawan in daftar_karyawan:
        jadwal_shift = [hitung_shift(karyawan, tanggal) for tanggal in rentang_tanggal]

        baris.append(
            [
                karyawan.kode_karyawan,
                karyawan.nama,
                *jadwal_shift,
            ]
        )

    return header, baris


def buat_csv(
    daftar_karyawan: Sequence[Karyawan],
    tanggal_mulai: date,
    tanggal_selesai: date,
) -> bytes:
    """Membuat file jadwal dalam format CSV."""

    header, baris = buat_matriks_jadwal(
        daftar_karyawan,
        tanggal_mulai,
        tanggal_selesai,
    )

    buffer_teks = StringIO(newline="")
    penulis = csv.writer(buffer_teks)

    penulis.writerow(header)
    penulis.writerows(baris)

    return buffer_teks.getvalue().encode("utf-8-sig")


def buat_excel(
    daftar_karyawan: Sequence[Karyawan],
    tanggal_mulai: date,
    tanggal_selesai: date,
) -> bytes:
    """Membuat file jadwal dalam format Excel."""

    header, baris = buat_matriks_jadwal(
        daftar_karyawan,
        tanggal_mulai,
        tanggal_selesai,
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Jadwal Shift"

    worksheet.append(header)

    for item in baris:
        worksheet.append(item)

    for sel in worksheet[1]:
        sel.font = Font(bold=True)
        sel.alignment = Alignment(horizontal="center")

    for baris_sel in worksheet.iter_rows(
        min_row=2,
        min_col=3,
    ):
        for sel in baris_sel:
            sel.alignment = Alignment(horizontal="center")

    worksheet.freeze_panes = "C2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.column_dimensions["A"].width = 12
    worksheet.column_dimensions["B"].width = 20

    for kolom in range(3, worksheet.max_column + 1):
        worksheet.column_dimensions[
            worksheet.cell(row=1, column=kolom).column_letter
        ].width = 13

    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()

    return buffer.getvalue()
