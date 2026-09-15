import csv
from io import BytesIO, StringIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook


def test_ekspor_csv_semua_karyawan(
    klien_api: TestClient,
) -> None:
    respons = klien_api.get(
        "/api/v1/jadwal/ekspor",
        params={
            "tanggal_mulai": "2024-12-26",
            "tanggal_selesai": "2024-12-29",
            "format": "csv",
        },
    )

    assert respons.status_code == 200
    assert respons.headers["content-type"].startswith("text/csv")

    assert (
        "jadwal-shift-2024-12-26-sampai-2024-12-29.csv"
        in respons.headers["content-disposition"]
    )

    teks = respons.content.decode("utf-8-sig")
    baris = list(csv.reader(StringIO(teks)))

    assert baris[0][0:2] == ["ID", "Nama"]
    assert len(baris) == 5


def test_ekspor_csv_satu_karyawan(
    klien_api: TestClient,
) -> None:
    respons = klien_api.get(
        "/api/v1/jadwal/ekspor",
        params={
            "kode_karyawan": "001",
            "tanggal_mulai": "2024-12-26",
            "tanggal_selesai": "2024-12-29",
            "format": "csv",
        },
    )

    assert respons.status_code == 200

    teks = respons.content.decode("utf-8-sig")
    baris = list(csv.reader(StringIO(teks)))

    assert len(baris) == 2
    assert baris[1][0] == "001"
    assert baris[1][1] == "Ahmad"


def test_ekspor_excel(
    klien_api: TestClient,
) -> None:
    respons = klien_api.get(
        "/api/v1/jadwal/ekspor",
        params={
            "tanggal_mulai": "2024-12-26",
            "tanggal_selesai": "2024-12-29",
            "format": "xlsx",
        },
    )

    assert respons.status_code == 200

    assert respons.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    workbook = load_workbook(BytesIO(respons.content))
    worksheet = workbook["Jadwal Shift"]

    assert worksheet["A1"].value == "ID"
    assert worksheet["B1"].value == "Nama"

    assert worksheet.max_row == 5
    assert worksheet.max_column == 6

    workbook.close()


def test_ekspor_karyawan_tidak_ditemukan(
    klien_api: TestClient,
) -> None:
    respons = klien_api.get(
        "/api/v1/jadwal/ekspor",
        params={
            "kode_karyawan": "999",
            "tanggal_mulai": "2025-01-01",
            "tanggal_selesai": "2025-01-05",
            "format": "csv",
        },
    )

    data = respons.json()

    assert respons.status_code == 404
    assert data["berhasil"] is False
    assert data["pesan"] == "Karyawan tidak ditemukan."


def test_ekspor_rentang_tanggal_tidak_valid(
    klien_api: TestClient,
) -> None:
    respons = klien_api.get(
        "/api/v1/jadwal/ekspor",
        params={
            "tanggal_mulai": "2025-01-10",
            "tanggal_selesai": "2025-01-01",
            "format": "csv",
        },
    )

    assert respons.status_code == 422


def test_format_ekspor_tidak_valid(
    klien_api: TestClient,
) -> None:
    respons = klien_api.get(
        "/api/v1/jadwal/ekspor",
        params={
            "tanggal_mulai": "2025-01-01",
            "tanggal_selesai": "2025-01-05",
            "format": "pdf",
        },
    )

    assert respons.status_code == 422
